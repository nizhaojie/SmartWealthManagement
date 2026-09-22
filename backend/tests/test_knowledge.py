import io
from collections.abc import Iterator
from datetime import datetime, timedelta, timezone

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session as OrmSession

from app.db.models import KnowledgeMeta
from app.knowledge.parsers import parse_document
from app.knowledge.seed_faq import FAQ_SOURCE_FILE, seed_faq
from app.knowledge.service import create_pending_document
from app.main import app
from app.settings import get_settings

EMPLOYEE_USERNAME = "advisor1"
SEEDED_PASSWORD = "Test@1234"


@pytest.fixture
def knowledge_client(auth_client: TestClient) -> Iterator[TestClient]:
    base_settings = get_settings()
    test_settings = base_settings.model_copy(
        update={
            "milvus_collection": base_settings.test_milvus_collection,
            "embedding_api_key": "",
        }
    )
    app.dependency_overrides[get_settings] = lambda: test_settings
    try:
        yield auth_client
    finally:
        app.dependency_overrides.pop(get_settings, None)


def _internal_token(client: TestClient) -> str:
    response = client.post(
        "/api/internal/auth/login",
        json={"username": EMPLOYEE_USERNAME, "password": SEEDED_PASSWORD},
    )
    return response.json()["data"]["access_token"]


def _auth_headers(client: TestClient) -> dict[str, str]:
    return {"Authorization": f"Bearer {_internal_token(client)}"}


def _upload(
    client: TestClient,
    *,
    filename: str,
    content: bytes,
    knowledge_type: str = "FAQ",
    content_type: str = "text/plain",
):
    return client.post(
        "/api/internal/knowledge/documents",
        headers=_auth_headers(client),
        files={"file": (filename, content, content_type)},
        data={"knowledge_type": knowledge_type},
    )


def _search(client: TestClient, *, query: str, knowledge_type: str | None = None, top_k: int = 5):
    body = {"query": query, "top_k": top_k}
    if knowledge_type is not None:
        body["knowledge_type"] = knowledge_type
    return client.post(
        "/api/internal/knowledge/search", headers=_auth_headers(client), json=body
    )


def _delete(client: TestClient, knowledge_id: int):
    return client.delete(
        f"/api/internal/knowledge/documents/{knowledge_id}", headers=_auth_headers(client)
    )


def _list(client: TestClient, *, knowledge_type: str | None = None, status: str | None = None):
    params = {}
    if knowledge_type is not None:
        params["knowledge_type"] = knowledge_type
    if status is not None:
        params["status"] = status
    return client.get(
        "/api/internal/knowledge/documents", headers=_auth_headers(client), params=params
    )


def _find(documents: list[dict], knowledge_id: int) -> dict:
    return next(doc for doc in documents if doc["knowledge_id"] == knowledge_id)


def test_upload_rejects_unsupported_format_with_business_error(knowledge_client):
    response = _upload(
        knowledge_client,
        filename="policy.pdf",
        content=b"%PDF-1.4 fake content",
        content_type="application/pdf",
    )

    assert response.status_code != 500
    body = response.json()
    assert body["code"] == 1006
    assert "格式" in body["message"]


def test_upload_returns_processing_immediately_then_becomes_active(knowledge_client):
    response = _upload(
        knowledge_client,
        filename="test_upload_txt.txt",
        content="随存随取型现金管理产品的赎回到账时间通常为下一个交易日。".encode("utf-8"),
    )

    assert response.status_code == 200
    data = response.json()["data"]
    assert data["status"] == "processing"
    assert data["knowledge_type"] == "FAQ"

    # TestClient 在 client.post() 返回前已经把 BackgroundTasks 跑完，
    # 所以这里立刻查列表就能看到处理结束后的终态，不需要真的轮询等待。
    listed = _find(_list(knowledge_client).json()["data"], data["knowledge_id"])
    assert listed["status"] == "active"
    assert listed["chunk_count"] >= 1

    _delete(knowledge_client, data["knowledge_id"])


def test_uploaded_document_is_retrievable_by_exact_chunk_text(knowledge_client):
    chunk_text = "随存随取型现金管理产品的赎回到账时间通常为下一个交易日。"
    upload_response = _upload(
        knowledge_client, filename="test_retrieve_exact.txt", content=chunk_text.encode("utf-8")
    )
    knowledge_id = upload_response.json()["data"]["knowledge_id"]

    try:
        search_response = _search(knowledge_client, query=chunk_text)
        assert search_response.status_code == 200
        payload = search_response.json()["data"]
        assert payload["score_threshold"] == get_settings().retrieval_score_threshold
        hits = payload["hits"]
        assert any(hit["knowledge_id"] == knowledge_id for hit in hits)
        matched = next(hit for hit in hits if hit["knowledge_id"] == knowledge_id)
        assert matched["content"] == chunk_text
        assert matched["source_file"] == "test_retrieve_exact.txt"
        assert matched["chunk_index"] == 0
        assert isinstance(matched["score"], float)
    finally:
        _delete(knowledge_client, knowledge_id)


def test_deleted_document_no_longer_appears_in_search_results(knowledge_client):
    chunk_text = "客户密码连续输错五次后账号将被临时锁定二十四小时。"
    upload_response = _upload(
        knowledge_client, filename="test_delete_then_search.txt", content=chunk_text.encode("utf-8")
    )
    knowledge_id = upload_response.json()["data"]["knowledge_id"]

    delete_response = _delete(knowledge_client, knowledge_id)
    assert delete_response.status_code == 200
    assert delete_response.json()["data"]["status"] == "expired"

    search_response = _search(knowledge_client, query=chunk_text)
    hits = search_response.json()["data"]["hits"]
    assert all(hit["knowledge_id"] != knowledge_id for hit in hits)


def test_deleting_unknown_document_returns_not_found(knowledge_client):
    response = _delete(knowledge_client, 999_999_999)
    assert response.status_code == 404


def test_markdown_heading_hierarchy_is_kept_as_chunk_metadata(knowledge_client):
    markdown = (
        "# 产品手册\n"
        "## 费率说明\n"
        "本产品的管理费率为百分之一点二每年，托管费率为百分之零点二每年。\n"
    )
    upload_response = _upload(
        knowledge_client,
        filename="test_markdown_heading.md",
        content=markdown.encode("utf-8"),
        knowledge_type="产品",
        content_type="text/markdown",
    )
    knowledge_id = upload_response.json()["data"]["knowledge_id"]

    try:
        search_response = _search(
            knowledge_client, query="本产品的管理费率为百分之一点二每年，托管费率为百分之零点二每年。"
        )
        hits = search_response.json()["data"]["hits"]
        matched = next(hit for hit in hits if hit["knowledge_id"] == knowledge_id)
        assert matched["heading_path"] == ["产品手册", "费率说明"]
    finally:
        _delete(knowledge_client, knowledge_id)


def test_docx_document_can_be_parsed_and_retrieved(knowledge_client):
    document = DocxDocument()
    document.add_heading("开户须知", level=1)
    document.add_paragraph("客户开户需提供本人身份证原件与银行卡，全程在网点办理。")
    buffer = io.BytesIO()
    document.save(buffer)

    upload_response = _upload(
        knowledge_client,
        filename="test_open_account.docx",
        content=buffer.getvalue(),
        knowledge_type="政策",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert upload_response.status_code == 200
    knowledge_id = upload_response.json()["data"]["knowledge_id"]

    try:
        search_response = _search(
            knowledge_client, query="客户开户需提供本人身份证原件与银行卡，全程在网点办理。"
        )
        hits = search_response.json()["data"]["hits"]
        matched = next(hit for hit in hits if hit["knowledge_id"] == knowledge_id)
        assert matched["heading_path"] == ["开户须知"]
    finally:
        _delete(knowledge_client, knowledge_id)


def test_search_can_be_scoped_to_a_knowledge_type(knowledge_client):
    faq_text = "客户经理更换流程需通过客服热线提交申请，三个工作日内完成。"
    policy_text = "客户经理更换事项适用公司内部人事调配政策的相关规定。"

    faq_upload = _upload(
        knowledge_client, filename="test_scope_faq.txt", content=faq_text.encode("utf-8"), knowledge_type="FAQ"
    )
    policy_upload = _upload(
        knowledge_client,
        filename="test_scope_policy.txt",
        content=policy_text.encode("utf-8"),
        knowledge_type="政策",
    )
    faq_id = faq_upload.json()["data"]["knowledge_id"]
    policy_id = policy_upload.json()["data"]["knowledge_id"]

    try:
        response = _search(knowledge_client, query=faq_text, knowledge_type="政策")
        hits = response.json()["data"]["hits"]
        assert all(hit["knowledge_id"] != faq_id for hit in hits)
    finally:
        _delete(knowledge_client, faq_id)
        _delete(knowledge_client, policy_id)


def test_faq_seed_imports_40_entries_and_they_are_retrievable(knowledge_client):
    base_settings = get_settings()
    seed_settings = base_settings.model_copy(
        update={
            "database_url": base_settings.test_database_url,
            "milvus_collection": base_settings.test_milvus_collection,
            "embedding_api_key": "",
        }
    )
    seed_faq(seed_settings)

    engine = create_engine(base_settings.test_database_url)
    try:
        with OrmSession(engine) as session:
            meta = session.scalar(
                select(KnowledgeMeta).where(KnowledgeMeta.source_file == FAQ_SOURCE_FILE)
            )
    finally:
        engine.dispose()

    assert meta is not None
    assert meta.status == "active"
    assert meta.chunk_count == 40

    response = _search(
        knowledge_client,
        query="密码连续输错五次后账号将被临时锁定二十四小时，此举用于防止密码被暴力破解。",
        knowledge_type="FAQ",
    )
    hits = response.json()["data"]["hits"]
    assert any(hit["knowledge_id"] == meta.id for hit in hits)


def test_fake_embedding_is_deterministic_across_repeated_searches(knowledge_client):
    chunk_text = "本产品最短持有期为九十天，提前赎回将收取百分之零点五的短期赎回费。"
    upload_response = _upload(
        knowledge_client, filename="test_determinism.txt", content=chunk_text.encode("utf-8")
    )
    knowledge_id = upload_response.json()["data"]["knowledge_id"]

    try:
        first = _search(knowledge_client, query=chunk_text).json()["data"]["hits"]
        second = _search(knowledge_client, query=chunk_text).json()["data"]["hits"]
        first_score = next(hit["score"] for hit in first if hit["knowledge_id"] == knowledge_id)
        second_score = next(hit["score"] for hit in second if hit["knowledge_id"] == knowledge_id)
        assert first_score == second_score
    finally:
        _delete(knowledge_client, knowledge_id)


def test_documents_list_returns_full_metadata_fields(knowledge_client):
    upload_response = _upload(
        knowledge_client,
        filename="test_list_fields.txt",
        content="本条记录用于校验文档列表返回的元数据字段是否齐全。".encode("utf-8"),
    )
    knowledge_id = upload_response.json()["data"]["knowledge_id"]

    try:
        listed = _find(_list(knowledge_client).json()["data"], knowledge_id)
        for field in (
            "knowledge_type",
            "title",
            "source_file",
            "version",
            "status",
            "create_time",
            "chunk_count",
            "expire_at",
        ):
            assert field in listed
        assert listed["source_file"] == "test_list_fields.txt"
        assert listed["version"] == "1"
        assert listed["status"] == "active"
    finally:
        _delete(knowledge_client, knowledge_id)


def test_documents_list_can_be_filtered_by_type_and_status(knowledge_client):
    faq_text = "文档列表筛选测试：FAQ 类型的一条示例内容。"
    policy_text = "文档列表筛选测试：政策类型的一条示例内容。"

    faq_id = _upload(
        knowledge_client, filename="test_filter_faq.txt", content=faq_text.encode("utf-8"), knowledge_type="FAQ"
    ).json()["data"]["knowledge_id"]
    policy_id = _upload(
        knowledge_client,
        filename="test_filter_policy.txt",
        content=policy_text.encode("utf-8"),
        knowledge_type="政策",
    ).json()["data"]["knowledge_id"]

    try:
        by_type = _list(knowledge_client, knowledge_type="政策").json()["data"]
        assert any(doc["knowledge_id"] == policy_id for doc in by_type)
        assert all(doc["knowledge_id"] != faq_id for doc in by_type)

        by_status = _list(knowledge_client, status="active").json()["data"]
        active_ids = {doc["knowledge_id"] for doc in by_status}
        assert {faq_id, policy_id} <= active_ids

        _delete(knowledge_client, faq_id)
        by_status_after_delete = _list(knowledge_client, status="expired").json()["data"]
        assert any(doc["knowledge_id"] == faq_id for doc in by_status_after_delete)
    finally:
        _delete(knowledge_client, policy_id)


def test_deleted_document_disappears_from_the_default_unfiltered_list(knowledge_client):
    knowledge_id = _upload(
        knowledge_client,
        filename="test_default_list_delete.txt",
        content="默认列表删除测试内容。".encode("utf-8"),
    ).json()["data"]["knowledge_id"]

    _delete(knowledge_client, knowledge_id)

    default_list = _list(knowledge_client).json()["data"]
    assert all(doc["knowledge_id"] != knowledge_id for doc in default_list)

    expired_only = _list(knowledge_client, status="expired").json()["data"]
    assert any(doc["knowledge_id"] == knowledge_id for doc in expired_only)


def test_create_pending_document_derives_expiry_from_the_injected_now(knowledge_client):
    # ADR-0011: 时间基准由调用方显式传入，所以这里能断言精确值，
    # 不需要像 HTTP 层的等价测试那样留误差容限。
    fixed_now = datetime(2026, 1, 1, 0, 0, 0)
    engine = create_engine(get_settings().test_database_url)
    try:
        with OrmSession(engine) as session:
            policy_meta = create_pending_document(
                session, filename="policy.txt", knowledge_type="政策", now=fixed_now
            )
            product_meta = create_pending_document(
                session, filename="product.txt", knowledge_type="产品", now=fixed_now
            )
            faq_meta = create_pending_document(
                session, filename="faq.txt", knowledge_type="FAQ", now=fixed_now
            )

            assert policy_meta.expire_at == fixed_now + timedelta(days=365)
            assert product_meta.expire_at == fixed_now + timedelta(days=180)
            assert faq_meta.expire_at is None

            # 这几条记录从未进入 ingest 流水线，没有 MinIO/Milvus 侧记录，
            # 直接删行即可，不需要走 delete_document 的三处联动。
            session.delete(policy_meta)
            session.delete(product_meta)
            session.delete(faq_meta)
            session.commit()
    finally:
        engine.dispose()


def test_policy_and_product_uploads_get_default_expiry_while_faq_does_not(knowledge_client):
    now = datetime.now(timezone.utc).replace(tzinfo=None)

    faq_id = _upload(
        knowledge_client,
        filename="test_expiry_faq.txt",
        content="过期时间默认值测试：FAQ。".encode("utf-8"),
        knowledge_type="FAQ",
    ).json()["data"]["knowledge_id"]
    product_id = _upload(
        knowledge_client,
        filename="test_expiry_product.txt",
        content="过期时间默认值测试：产品。".encode("utf-8"),
        knowledge_type="产品",
    ).json()["data"]["knowledge_id"]
    policy_id = _upload(
        knowledge_client,
        filename="test_expiry_policy.txt",
        content="过期时间默认值测试：政策。".encode("utf-8"),
        knowledge_type="政策",
    ).json()["data"]["knowledge_id"]

    try:
        documents = _list(knowledge_client).json()["data"]
        faq_doc = _find(documents, faq_id)
        product_doc = _find(documents, product_id)
        policy_doc = _find(documents, policy_id)

        assert faq_doc["expire_at"] is None

        product_expire_at = datetime.fromisoformat(product_doc["expire_at"])
        policy_expire_at = datetime.fromisoformat(policy_doc["expire_at"])
        assert abs((product_expire_at - (now + timedelta(days=180))).total_seconds()) < 300
        assert abs((policy_expire_at - (now + timedelta(days=365))).total_seconds()) < 300
    finally:
        _delete(knowledge_client, faq_id)
        _delete(knowledge_client, product_id)
        _delete(knowledge_client, policy_id)


def test_corrupt_docx_upload_is_marked_failed_with_parse_stage_and_reason(knowledge_client):
    upload_response = _upload(
        knowledge_client,
        filename="test_corrupt.docx",
        content=b"this is not a real docx file",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    knowledge_id = upload_response.json()["data"]["knowledge_id"]

    try:
        listed = _find(_list(knowledge_client).json()["data"], knowledge_id)
        assert listed["status"] == "failed"
        assert listed["stage"] == "parse"
        assert listed["failure_reason"]
    finally:
        _delete(knowledge_client, knowledge_id)


# --- FAQ 按问答对拆分（`rag-retrieval-upgrade` 01）---


QA_PAIRS_TXT = (
    "公司的客服电话是多少?\t400-XXX-XXXX，服务时间 7:00-22:00。\n"
    "基金申购后多久确认?\t交易日 15:00 前提交的申请 T+1 日确认份额。\n"
    "什么是七日年化收益率?\t过去七天每万份基金份额净收益折合成的年收益率。\n"
)


def test_txt_qa_pairs_become_one_section_each():
    """一组问答是 FAQ 的最小完整语义单元：一行一组，一组一节。"""
    sections = parse_document("faq.txt", QA_PAIRS_TXT.encode("utf-8"))

    assert [section.heading_path for section in sections] == [
        ["公司的客服电话是多少?"],
        ["基金申购后多久确认?"],
        ["什么是七日年化收益率?"],
    ]
    assert [section.text for section in sections] == [
        "公司的客服电话是多少?\n400-XXX-XXXX，服务时间 7:00-22:00。",
        "基金申购后多久确认?\n交易日 15:00 前提交的申请 T+1 日确认份额。",
        "什么是七日年化收益率?\n过去七天每万份基金份额净收益折合成的年收益率。",
    ]


def test_txt_answer_keeps_the_tabs_after_the_first_separator():
    """只按第一个 tab 切分，答案内部再有 tab 不二次切分。"""
    content = "问题?\t答案第一段\t答案第二段\n".encode("utf-8")

    sections = parse_document("faq.txt", content)

    assert [(section.heading_path, section.text) for section in sections] == [
        (["问题?"], "问题?\n答案第一段\t答案第二段")
    ]


def test_txt_plain_lines_and_qa_pairs_keep_the_original_order():
    """不含 tab 的行按原样成节（`heading_path` 为空），与问答节按出现顺序交错。"""
    content = "前言第一行。\n问?\t答。\n说明第二行。\n".encode("utf-8")

    sections = parse_document("mixed.txt", content)

    assert [section.heading_path for section in sections] == [[], ["问?"], []]
    assert [section.text for section in sections] == ["前言第一行。", "问?\n答。", "说明第二行。"]


def test_txt_without_tabs_is_one_section_like_before():
    """没有 tab 的 txt 与改动前同形：整篇一个 Section，正文是去首尾空白的原文。"""
    content = "第一段没有制表符。\n\n第二段也没有。\n".encode("utf-8")

    sections = parse_document("plain.txt", content)

    assert len(sections) == 1
    assert sections[0].heading_path == []
    assert sections[0].text == content.decode("utf-8").strip()


def test_qa_pair_txt_is_chunked_one_pair_per_chunk(knowledge_client):
    upload_response = _upload(
        knowledge_client,
        filename="test_qa_pairs.txt",
        content=QA_PAIRS_TXT.encode("utf-8"),
        knowledge_type="FAQ",
    )
    knowledge_id = upload_response.json()["data"]["knowledge_id"]

    try:
        listed = _find(_list(knowledge_client).json()["data"], knowledge_id)
        # 一组问答远短于 512 token，块数因此等于问答对数，不是滑动窗口切出来的块数。
        assert listed["chunk_count"] == 3

        for line in QA_PAIRS_TXT.strip().split("\n"):
            question, answer = line.split("\t", 1)
            hits = _search(
                knowledge_client, query=f"{question}\n{answer}", knowledge_type="FAQ"
            ).json()["data"]["hits"]
            matched = next(hit for hit in hits if hit["knowledge_id"] == knowledge_id)
            assert matched["content"] == f"{question}\n{answer}"
            assert matched["heading_path"] == [question]
    finally:
        _delete(knowledge_client, knowledge_id)
